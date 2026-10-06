"""Excel access — Kasus Perselisihan.

Same satudata source as the other Kemnaker datasets; only the workbook layout
differs. The workbook has one sheet per month (Januari..Desember) plus a
``metadata`` sheet. Each month sheet has a three-row header (Excel rows 4-6)
covering Jenis Perselisihan, Jumlah Perselisihan (Pusat/Daerah/Jumlah),
Penyelesaian Perselisihan (Bipartit/Mediasi PB/Mediasi Anjuran/Konsiliasi/
Arbitrasi/Hakim HI/Jumlah) and Sisa Kasus.
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any, Dict, List

import openpyxl

from .._satudata import SatuDataExcelExtractor

DATA_START = 6                # 0-based Excel row 7 (first data row)
TITLE_ROW = 1                 # 0-based Excel row 2 -> 'Tahun <yyyy>'
# Column index -> model field name (0-based).
FIELD_BY_COLUMN = {
    0: "jenis_perselisihan",
    1: "jumlah_pusat",
    2: "jumlah_daerah",
    3: "jumlah",
    4: "bipartit",
    5: "mediasi_pb",
    6: "mediasi_anjuran",
    7: "konsiliasi",
    8: "arbitrasi",
    9: "hakim_hi",
    10: "jumlah_penyelesaian",
    11: "sisa_kasus",
}
MONTH_NUMBER = {
    "januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5, "juni": 6,
    "juli": 7, "agustus": 8, "september": 9, "oktober": 10, "november": 11,
    "desember": 12,
}


class Extractor(SatuDataExcelExtractor):
    """Resolve the perselisihan workbook URL from the satudata API, then parse it."""

    def __init__(self, config: Dict[str, Any]):
        merged = dict(config)
        merged.setdefault("require_columns", "jenis_perselisihan")
        merged.setdefault("skip_row_prefixes", "total,sumber")
        super().__init__(merged)

    def _parse_excel(self, content: bytes) -> List[Dict[str, Any]]:
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
        records: List[Dict[str, Any]] = []
        for name in workbook.sheetnames:
            month = MONTH_NUMBER.get(name.strip().casefold())
            if not month:                       # skips the 'metadata' sheet
                continue
            rows = list(workbook[name].iter_rows(values_only=True))
            tahun = self._year(rows)
            periode = f"{tahun:04d}-{month:02d}" if tahun else None
            for row in rows[DATA_START:]:
                record = {
                    field: self._clean(row[column] if column < len(row) else None)
                    for column, field in FIELD_BY_COLUMN.items()
                }
                if not self._is_data_row(record):
                    continue
                record["jenis_perselisihan"] = _canonical_jenis(record["jenis_perselisihan"])
                record["periode"] = periode
                record["tahun"] = tahun
                records.append(record)
        return records

    @staticmethod
    def _clean(value: Any) -> Any:
        if isinstance(value, str):
            return " ".join(value.split()) or None
        return value

    @staticmethod
    def _year(rows: List[tuple]) -> int | None:
        if len(rows) <= TITLE_ROW or not rows[TITLE_ROW]:
            return None
        match = re.search(r"(\d{4})", str(rows[TITLE_ROW][0]))
        return int(match.group(1)) if match else None


def _canonical_jenis(value: Any) -> str | None:
    """Normalize the 'Perselisihan ...' labels to a short canonical form."""
    text = " ".join(str(value).split()) if value is not None else ""
    low = text.casefold()
    if "phk" in low:
        return "PHK"
    if "serikat" in low:
        return "Serikat Pekerja"
    return text or None
