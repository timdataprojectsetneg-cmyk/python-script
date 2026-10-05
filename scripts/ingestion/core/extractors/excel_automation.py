#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
excel_automation.py
===================
Otomasi membaca file Excel (sheet data + sheet "Standar Data & Metadata") dan
mengubahnya menjadi sekumpulan tabel terstruktur. Keluarannya CSV, bukan load ke
Postgres, sehingga bisa dipakai ulang di dalam usecase: hasil CSV menjadi input
untuk proses normalisasi dan loaders.

Alur:
  1. Baca semua sheet                     -> {prefix}_meta_sheet
  2. Temukan header (multi-level + merge) -> otomatis
  3. Kenali kolom (label vs ukuran)       -> {prefix}_meta_column
  4. Deteksi tipe data                    -> {prefix}_data, {prefix}_data_long
  5. Pahami arti kolom (dari metadata)    -> {prefix}_dim_mapping + deskripsi di meta_column
  6. Buat metadata tabular                -> {prefix}_meta_dataset, {prefix}_meta_indicator
  7. Profiling (opsional)                 -> {prefix}_profile
  8. Relasi antar-kolom (opsional)        -> {prefix}_relation
  9. Rekomendasi                          -> {prefix}_recommendation

Pakai sebagai library di dalam usecase:
    from ingestion.core.extractors.excel_automation import (
        extract_workbook, write_csvs,
    )

    result = extract_workbook("file.xlsx", const={"provinsi": "Gorontalo"})
    write_csvs(result.tables, out_dir="./out")

Pakai sebagai CLI (hasil selalu CSV):
    python -m ingestion.core.extractors.excel_automation file.xlsx --out-dir ./out
    python -m ingestion.core.extractors.excel_automation file.xlsx --output data_long
    python -m ingestion.core.extractors.excel_automation file.xlsx --output data,data_long,recommendation
    python -m ingestion.core.extractors.excel_automation file.xlsx --output all
    # pilihan: data, data_long, dim_mapping, meta_sheet, meta_dataset, meta_indicator,
    #          meta_column, profile, relation, recommendation, all
"""
from __future__ import annotations

import argparse
import difflib
import itertools
import os
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# =============================================================================
# KONFIGURASI (ubah di sini bila format file berbeda)
# =============================================================================
MAX_IDENT = 55                      # batas panjang nama kolom (Postgres max 63 byte)
DASH_AS_ZERO = True                 # sel '-' dianggap 0 (konvensi BPS). False -> NULL
TOTAL_LABELS = {"total", "jumlah", "jumlah total", "grand total"}
METADATA_SHEET_HINTS = ("metadata", "standar data", "kamus", "dictionary")
NULL_TOKENS = {"n/a", "na", "nan", "null", "none", "..", "..."}
LEVEL_NAMES = ["", "kategori", "subkategori"]      # nama kolom label bertingkat
MATCH_THRESHOLD = 0.6               # ambang kemiripan nama dimensi <-> metadata
STOPWORDS = {
    "yang", "di", "dan", "atau", "seluruh", "kategori", "klasifikasi", "jumlah",
    "seminggu", "lalu", "tahun", "dari", "pada", "untuk", "dengan", "ke", "menurut",
}
# Singkatan/istilah yang tidak mirip secara teks dengan nama klasifikasi di metadata.
# Tambahkan sesuai domain data Anda.
SYNONYMS = {
    "kbji": "jabatan",
    "kbli": "lapangan usaha",
    "perkotaan": "daerah tempat tinggal",
    "perdesaan": "daerah tempat tinggal",
    "pedesaan": "daerah tempat tinggal",
}
MONTHS = ("Januari|Februari|Maret|April|Mei|Juni|Juli|Agustus|September|Oktober|"
          "Nopember|November|Desember")

# =============================================================================
# UTILITAS DASAR
# =============================================================================
RE_ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
RE_ISO_TS = re.compile(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2})?")
_RE_ID = re.compile(r"^[-+]?\d{1,3}(\.\d{3})+(,\d+)?$")     # 1.234.567,89
_RE_EN = re.compile(r"^[-+]?\d{1,3}(,\d{3})+(\.\d+)?$")     # 1,234,567.89
_RE_PLAIN = re.compile(r"^[-+]?\d+([.,]\d+)?$")             # 12 / 12,5 / 12.5
KV_RE = re.compile(r"^\s*([^:\n]{1,40}?)\s*:\s*(\S.*)$", re.S)


class _Sentinel:
    def __init__(self, n): self.n = n
    def __repr__(self): return self.n


DASH = _Sentinel("DASH")          # sel berisi '-' / ' -   '
NOT_NUM = _Sentinel("NOT_NUM")    # string yang bukan angka


def is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.replace("\xa0", " ").strip() == "")


def clean_text(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def one_line(v: Any) -> str:
    """Baris pertama sel (header sering berisi catatan di baris berikutnya)."""
    return clean_text(str(v).replace("\xa0", " ").split("\n")[0])


def slugify(text: Any, maxlen: int = MAX_IDENT) -> str:
    s = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-zA-Z0-9]+", "_", s).strip("_").lower() or "col"
    if s[0].isdigit():
        s = "c_" + s
    return s[:maxlen].rstrip("_")


def shorten_name(parts: list[str], maxlen: int = MAX_IDENT) -> str:
    """Slug dari path header; bila terlalu panjang buang level paling atas satu per satu."""
    for k in range(len(parts)):
        cand = slugify(" ".join(parts[k:]), maxlen=200)
        if len(cand) <= maxlen:
            return cand
    return slugify(parts[-1], maxlen) if parts else "col"


def dedupe(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for n in names:
        if n in seen:
            seen[n] += 1
            n = f"{n[:MAX_IDENT - 4]}_{seen[n]}"
        else:
            seen[n] = 1
        out.append(n)
    # pastikan unik setelah suffix
    final, used = [], set()
    for n in out:
        base, i = n, 2
        while n in used:
            n = f"{base}_{i}"; i += 1
        used.add(n); final.append(n)
    return final


def parse_number(v: Any):
    """Return int/float, None (kosong), DASH, atau NOT_NUM."""
    if v is None:
        return None
    if isinstance(v, bool):
        return NOT_NUM
    if isinstance(v, (int, np.integer)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        if np.isnan(v):
            return None
        return int(v) if float(v).is_integer() else float(v)
    if isinstance(v, (datetime, date)):
        return NOT_NUM
    s = str(v).replace("\xa0", " ").strip()
    if s == "":
        return None
    if re.fullmatch(r"[-–—\s]+", s):
        return DASH
    if s.casefold() in NULL_TOKENS:
        return None
    s2 = s.replace(" ", "")
    if _RE_ID.match(s2):
        s2 = s2.replace(".", "").replace(",", ".")
    elif _RE_EN.match(s2):
        s2 = s2.replace(",", "")
    elif _RE_PLAIN.match(s2):
        s2 = s2.replace(",", ".")
    else:
        return NOT_NUM
    try:
        f = float(s2)
    except ValueError:
        return NOT_NUM
    return int(f) if (f.is_integer() and "." not in s2) else f


def is_numlike(v: Any) -> bool:
    r = parse_number(v)
    return r is not None and r is not NOT_NUM


def infer_dtype(values) -> str:
    vals = [v for v in values if v is not None and not (isinstance(v, float) and np.isnan(v))]
    if not vals:
        return "text"
    if all(isinstance(v, bool) for v in vals):
        return "boolean"
    if any(isinstance(v, bool) for v in vals):
        return "text"
    if all(isinstance(v, (datetime, date)) for v in vals):
        has_time = any(isinstance(v, datetime) and (v.hour or v.minute or v.second) for v in vals)
        return "timestamp" if has_time else "date"
    if all(isinstance(v, (int, np.integer)) for v in vals):
        return "bigint"
    if all(isinstance(v, (int, float, np.integer, np.floating)) for v in vals):
        return "numeric"
    if all(isinstance(v, str) for v in vals):
        if all(RE_ISO_DATE.fullmatch(v.strip()) for v in vals):
            return "date"
        if all(RE_ISO_TS.fullmatch(v.strip()) for v in vals):
            return "timestamp"
    return "text"


def _null(v) -> bool:
    if v is None or v is pd.NA or v is pd.NaT:
        return True
    return isinstance(v, (float, np.floating)) and bool(np.isnan(v))


def to_series(values, dtype: str):
    values = [None if _null(v) else v for v in values]
    if dtype == "bigint":
        return pd.array([None if v is None else int(v) for v in values], dtype="Int64")
    if dtype == "numeric":
        return pd.Series([np.nan if v is None else float(v) for v in values], dtype="float64")
    if dtype == "boolean":
        return pd.array([None if v is None else bool(v) for v in values], dtype="boolean")
    if dtype == "date":
        s = pd.to_datetime(pd.Series(values), errors="coerce")
        return pd.Series([None if pd.isna(x) else x.date() for x in s], dtype="object")
    if dtype == "timestamp":
        return pd.to_datetime(pd.Series(values), errors="coerce")
    return pd.Series([None if v is None else str(v) for v in values], dtype="object")


def kv_from_text(text: Any):
    m = KV_RE.match(str(text))
    return (m.group(1).strip(), m.group(2).strip()) if m else None


# =============================================================================
# 1. BACA SEMUA SHEET
# =============================================================================
@dataclass
class Grid:
    name: str
    values: list
    merged: list            # (r1, c1, r2, c2) 0-based inclusive

    @property
    def n_rows(self): return len(self.values)

    @property
    def n_cols(self): return len(self.values[0]) if self.values else 0

    def filled(self):
        """Salinan grid dengan sel merge diisi nilai sel kiri-atas."""
        out = [row[:] for row in self.values]
        for r1, c1, r2, c2 in self.merged:
            if r1 >= self.n_rows or c1 >= self.n_cols:
                continue
            v = self.values[r1][c1]
            for r in range(r1, min(r2, self.n_rows - 1) + 1):
                for c in range(c1, min(c2, self.n_cols - 1) + 1):
                    out[r][c] = v
        return out


def load_grids(path: str) -> list[Grid]:
    wb_f = load_workbook(path, data_only=False)     # untuk membaca formula
    wb_v = load_workbook(path, data_only=True)      # nilai hasil hitung (cache)
    grids = []
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        nr, nc = ws_f.max_row, ws_f.max_column
        vals = [[ws_v.cell(r, c).value for c in range(1, nc + 1)] for r in range(1, nr + 1)]
        # Formula tanpa cache nilai (mis. '=C26') -> selesaikan referensi sederhana
        for r in range(nr):
            for c in range(nc):
                if vals[r][c] is None:
                    f = ws_f.cell(r + 1, c + 1).value
                    if isinstance(f, str) and f.startswith("="):
                        m = re.fullmatch(r"=\$?([A-Z]{1,3})\$?(\d+)", f.strip())
                        if m:
                            vals[r][c] = ws_v[f"{m.group(1)}{m.group(2)}"].value
        # pangkas baris/kolom kosong di ujung
        last_r = max((r for r in range(nr) if any(not is_blank(x) for x in vals[r])), default=-1)
        last_c = max((c for c in range(nc) if any(not is_blank(vals[r][c]) for r in range(nr))), default=-1)
        vals = [row[:last_c + 1] for row in vals[:last_r + 1]]
        merged = [(rg.min_row - 1, rg.min_col - 1, rg.max_row - 1, rg.max_col - 1)
                  for rg in ws_f.merged_cells.ranges]
        grids.append(Grid(ws_f.title, vals, merged))
    return grids


def detect_roles(grids: list[Grid]) -> dict[str, str]:
    roles = {g.name: ("metadata" if any(h in g.name.casefold() for h in METADATA_SHEET_HINTS) else "data")
             for g in grids}
    if "metadata" not in roles.values() and len(grids) > 1:      # tidak ada petunjuk nama
        for g in grids[1:]:
            roles[g.name] = "metadata"
    if "data" not in roles.values():
        roles[grids[0].name] = "data"
    return roles


# =============================================================================
# 2-3. HEADER, KOLOM, PATH HEADER MULTI-LEVEL
# =============================================================================
def header_paths(filled: list, hs: int, ds: int, n_cols: int) -> list[list[str]]:
    """Path header per kolom, mis. ['Jenis Kelamin', 'Laki-laki']. Menangani merge
    maupun 'center across selection' (sel kosong di kanan header induk)."""
    H = [[filled[r][c] if c < len(filled[r]) else None for c in range(n_cols)] for r in range(hs, ds)]
    for i in range(len(H)):
        for c in range(1, n_cols):
            if is_blank(H[i][c]) and not is_blank(H[i][c - 1]):
                if any(not is_blank(H[j][c]) for j in range(i + 1, len(H))):
                    H[i][c] = H[i][c - 1]
    paths = []
    for c in range(n_cols):
        parts: list[str] = []
        for i in range(len(H)):
            if not is_blank(H[i][c]):
                t = one_line(H[i][c])
                if not parts or parts[-1] != t:
                    parts.append(t)
        paths.append(parts)
    return paths


@dataclass
class Layout:
    header_start: int
    data_start: int
    data_end: int
    n_cols: int
    body_rows: list
    title_rows: list = field(default_factory=list)
    footer_rows: list = field(default_factory=list)
    dropped_rows: list = field(default_factory=list)


def detect_layout(g: Grid, data_start_override: Optional[int] = None) -> Layout:
    stats = []
    for r in range(g.n_rows):
        nb = [v for v in g.values[r] if not is_blank(v)]
        stats.append((len(nb), sum(1 for v in nb if is_numlike(v))))
    if data_start_override:
        ds = data_start_override - 1
    else:
        ds = next((r for r, (nb, nu) in enumerate(stats) if nb > 0 and nu >= 1 and nu / nb >= 0.5), None)
        if ds is None:
            raise ValueError(f"Sheet '{g.name}': tidak ditemukan baris data numerik")
    hs = next((r for r in range(ds) if stats[r][0] >= 2), ds)
    data_end = max(r for r in range(g.n_rows) if stats[r][1] >= 1)
    body, dropped = [], []
    for r in range(ds, data_end + 1):
        if stats[r][0] == 0:
            continue
        (body if stats[r][1] >= 1 else dropped).append(r)
    footer = [r for r in range(data_end + 1, g.n_rows) if stats[r][0] > 0]
    title = [r for r in range(hs) if stats[r][0] > 0]
    n_cols = max(c for r in range(hs, data_end + 1) for c in range(g.n_cols) if not is_blank(g.values[r][c])) + 1
    return Layout(hs, ds, data_end, n_cols, body, title, footer, dropped)


def name_columns(paths, roles, letters) -> list[str]:
    n = len(paths)
    names: list[Optional[str]] = [None] * n
    i = 0
    while i < n:
        if roles[i] == "label":
            j = i
            while j + 1 < n and roles[j + 1] == "label" and paths[j + 1] == paths[i]:
                j += 1
            base = shorten_name(paths[i]) if paths[i] else f"col_{letters[i].lower()}"
            for k, c in enumerate(range(i, j + 1)):
                names[c] = base if k == 0 else (LEVEL_NAMES[k] if k < len(LEVEL_NAMES) else f"level_{k + 1}")
            i = j + 1
        else:
            names[i] = shorten_name(paths[i]) if paths[i] else f"col_{letters[i].lower()}"
            i += 1
    return dedupe([n_ or "col" for n_ in names])


# =============================================================================
# 5-6. METADATA (STANDAR DATA & METADATA)
# =============================================================================
@dataclass
class Metadata:
    dataset: pd.DataFrame
    indicator: Optional[pd.DataFrame]
    indicator_paths: dict
    indicator_letters: dict
    klasifikasi: list
    info: dict


def parse_kv_block(g: Grid, end_row: int) -> list[dict]:
    rows, section = [], None

    def add(key, value, r, c):
        rows.append(dict(sheet_name=g.name, section=section, key=key, value=value,
                         source_cell=f"{get_column_letter(c + 1)}{r + 1}"))

    for r in range(end_row):
        cells = [(c, v) for c, v in enumerate(g.values[r]) if not is_blank(v)]
        if not cells:
            continue
        c0, v0 = cells[0]
        if len(cells) == 1 and c0 == 0 and isinstance(v0, str) and not kv_from_text(v0):
            section = clean_text(v0)
            continue
        consumed = set()
        if c0 == 0 and isinstance(v0, str) and not kv_from_text(v0):
            label, val = clean_text(v0), None
            for c, v in cells[1:]:
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    continue                                   # nomor urut kolom
                if isinstance(v, str) and kv_from_text(v):
                    break
                val = (c, v)
                break
            consumed.add(0)
            if val is not None:
                add(label, str(val[1]).strip(), r, val[0]); consumed.add(val[0])
            elif not label.startswith("("):
                add(label, None, r, 0)
        for c, v in cells:
            if c in consumed or not isinstance(v, str):
                continue
            kv = kv_from_text(v)
            if kv:
                add(kv[0], kv[1], r, c)
    return rows


def parse_metadata(g: Grid) -> tuple[Metadata, dict]:
    filled = g.filled()
    # --- cari tabel indikator: baris yang memuat sel 'No.' ---
    hr = None
    for r in range(g.n_rows):
        txt = [clean_text(v).casefold() for v in g.values[r] if isinstance(v, str)]
        if any(t in ("no.", "no") for t in txt) and sum(not is_blank(v) for v in g.values[r]) >= 4:
            hr = r
            break
    kv_rows = parse_kv_block(g, hr if hr is not None else g.n_rows)

    indicator, ind_paths, ind_letters, ind_info = None, {}, {}, {}
    if hr is not None:
        code_r = None
        for r in range(hr + 1, min(hr + 7, g.n_rows)):
            nb = [v for v in g.values[r] if not is_blank(v)]
            if len(nb) >= 4 and all(isinstance(v, (int, float)) and not isinstance(v, bool)
                                    and float(v).is_integer() for v in nb):
                code_r = r
                break
        ds = (code_r + 1) if code_r is not None else hr + 1
        n_cols = g.n_cols
        paths = header_paths(filled, hr, code_r if code_r is not None else hr + 1, n_cols)
        names = dedupe([shorten_name(p) if p else f"col_{get_column_letter(i + 1).lower()}"
                        for i, p in enumerate(paths)])
        data = []
        for r in range(ds, g.n_rows):
            row = g.values[r]
            if all(is_blank(v) for v in row):
                break
            data.append(row)
        cols = {}
        for i, nm in enumerate(names):
            vals = [None if is_blank(row[i]) else (row[i].strip() if isinstance(row[i], str) else row[i])
                    for row in data]
            dt = infer_dtype(vals)
            cols[nm] = to_series([str(v) if (dt == "text" and v is not None) else v for v in vals], dt)
        indicator = pd.DataFrame(cols)
        ind_paths = {nm: paths[i] for i, nm in enumerate(names)}
        ind_letters = {nm: get_column_letter(i + 1) for i, nm in enumerate(names)}
        ind_info = dict(header_first=hr + 1, header_last=(code_r if code_r is not None else hr + 1),
                        data_first=ds + 1, data_last=ds + len(data), n_indicators=len(data))

    ds_df = pd.DataFrame(kv_rows)
    if ds_df.empty:
        ds_df = pd.DataFrame(columns=["sheet_name", "section", "key", "value", "source_cell"])
    ds_df.insert(len(ds_df.columns), "ordinal", range(1, len(ds_df) + 1))

    # daftar klasifikasi (dari kunci 'Klasifikasi' dan kolom klasifikasi di tabel indikator)
    texts = [str(v) for k, v in zip(ds_df["key"], ds_df["value"]) if str(k).casefold() == "klasifikasi" and v]
    if indicator is not None and "klasifikasi" in indicator.columns:
        texts += [str(v) for v in indicator["klasifikasi"].dropna()]
    klas, seen = [], set()
    for t in texts:
        for line in t.split("\n"):
            item = re.sub(r"^\s*\d+[.)]\s*", "", line).strip()
            if item and item.casefold() not in seen:
                seen.add(item.casefold()); klas.append(item)

    info = dict(ind_info, n_kv=len(ds_df), n_klasifikasi=len(klas))
    return Metadata(ds_df, indicator, ind_paths, ind_letters, klas, info), info


def indicator_info(meta: Metadata, col_path: list[str]) -> dict:
    out = dict(nama=None, definisi=None, ukuran=None, satuan=None, interpretasi=None, source=None)
    ind = meta.indicator

    def pick(row, *keys):
        for c in ind.columns:
            if any(c == k or c.startswith(k) for k in keys):
                v = row[c]
                return None if pd.isna(v) else str(v)
        return None

    if ind is not None and len(ind):
        row = None
        if len(ind) == 1:
            row = ind.iloc[0]
        else:
            best, bs = None, 0.0
            for _, rw in ind.iterrows():
                s = text_score(" ".join(col_path), pick(rw, "nama_indikator") or "")[0]
                if s > bs:
                    best, bs = rw, s
            if bs >= MATCH_THRESHOLD:
                row = best
        if row is not None:
            out.update(nama=pick(row, "nama_indikator"), definisi=pick(row, "definisi"),
                       ukuran=pick(row, "ukuran"), satuan=pick(row, "satuan"),
                       interpretasi=pick(row, "interpretasi"),
                       source="metadata: tabel indikator")
            return out
    kv = {str(k).casefold(): v for k, v in zip(meta.dataset["key"], meta.dataset["value"]) if v}
    if kv:
        out.update(nama=kv.get("judul data") or kv.get("konsep"), definisi=kv.get("definisi"),
                   ukuran=kv.get("ukuran"), satuan=kv.get("satuan"),
                   source="metadata: blok Standar Data")
    return out


def describe_measure(info: dict, path: list[str]) -> tuple[Optional[str], Optional[str], Optional[str]]:
    if not any(info.get(k) for k in ("nama", "ukuran", "satuan")):
        return None, None, None
    parts = [info["nama"]] if info.get("nama") else []
    if path:
        parts.append("kolom: " + " > ".join(path))
    if info.get("ukuran"):
        parts.append(f"ukuran: {info['ukuran']}")
    if info.get("satuan"):
        parts.append(f"satuan: {info['satuan']}")
    return " | ".join(parts), info.get("definisi"), info.get("source")


# --- pencocokan nama dimensi di data <-> daftar klasifikasi di metadata ---
def _norm(text: str) -> str:
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()


def tokens(text: str) -> set:
    toks = set()
    for w in re.findall(r"[a-z]+", _norm(text)):
        if w in STOPWORDS:
            continue
        toks.add(w)
        if w in SYNONYMS:
            toks |= set(SYNONYMS[w].split())
    return toks


def text_score(src: str, cand: str) -> tuple[float, str]:
    a, b = tokens(src), tokens(cand)
    if not a or not b:
        return 0.0, "none"
    if " ".join(sorted(a)) == " ".join(sorted(b)):
        return 1.0, "exact"
    inter = len(a & b)
    contain, jacc = inter / len(b), inter / len(a | b)
    seq = difflib.SequenceMatcher(None, _norm(src), _norm(cand)).ratio()
    return max(0.7 * contain + 0.3 * jacc, seq), "token_overlap"


def match_dimension(src: str, klasifikasi: list[str]) -> tuple[Optional[str], float, str]:
    best, bs, bm = None, 0.0, "none"
    for k in klasifikasi:
        s, m = text_score(src, k)
        if s > bs:
            best, bs, bm = k, s, m
    if bs < MATCH_THRESHOLD:
        return None, round(bs, 3), "none"
    return best, round(bs, 3), bm


# =============================================================================
# 7-8. PROFILING & RELASI
# =============================================================================
def profile_df(sheet, table, df, cols, dtypes, extra) -> pd.DataFrame:
    rows = []
    for c in cols:
        s = df[c]
        n = len(s)
        nn = int(s.isna().sum())
        d = dict(sheet_name=sheet, table_name=table, column_name=c, inferred_type=dtypes[c],
                 n_rows=n, n_null=nn, null_pct=round(100 * nn / n, 2) if n else None,
                 n_distinct=int(s.nunique(dropna=True)))
        nz = s.dropna()
        if dtypes[c] in ("bigint", "numeric") and len(nz):
            f = nz.astype(float)
            d.update(min_value=str(nz.min()), max_value=str(nz.max()), mean_value=f.mean(),
                     median_value=f.median(), std_value=f.std(), sum_value=f.sum())
        elif len(nz):
            st = nz.astype(str)
            vc = st.value_counts()
            d.update(min_value=st.min(), max_value=st.max(), top_value=vc.index[0], top_freq=int(vc.iloc[0]),
                     min_len=int(st.str.len().min()), max_len=int(st.str.len().max()))
        d.update(extra.get(c, {}))
        rows.append(d)
    out = pd.DataFrame(rows)
    for c in ("mean_value", "median_value", "std_value", "sum_value"):
        if c not in out:
            out[c] = np.nan
        out[c] = pd.to_numeric(out[c])
    for c in ("top_freq", "min_len", "max_len", "dash_cells"):
        if c not in out:
            out[c] = None
        out[c] = pd.array(out[c], dtype="Int64")
    return out


def analyze_relations(sheet, df, label_cols, measure_cols, paths, is_total_row_col="is_total_row"):
    rel = []

    def add(rtype, left, right, expr, checked=None, passed=None, failed=None, status="info",
            detail=None, value=None):
        rel.append(dict(sheet_name=sheet, relation_type=rtype, left_scope=left, right_scope=right,
                        expression=expr, checked=checked, passed=passed, failed=failed,
                        status=status, detail=detail, value=value))

    n = len(df)
    # R1. kandidat kunci
    key = None
    for k in range(1, len(label_cols) + 1):
        for combo in itertools.combinations(label_cols, k):
            sub = df[list(combo)]
            if not sub.isna().any().any() and not sub.duplicated().any():
                key = list(combo)
                break
        if key:
            break
    if key:
        add("candidate_key", ", ".join(key), "(seluruh kolom)", f"UNIQUE({', '.join(key)})", n, n, 0, "ok",
            "Kombinasi unik dan tidak ada NULL -> layak jadi primary key")
    else:
        add("candidate_key", ", ".join(label_cols), "(seluruh kolom)", "UNIQUE(label_cols)", n, None, None,
            "warn", "Tidak ada kombinasi kolom label yang unik")

    # R2. hierarki antar kolom label
    for p, c in zip(label_cols[:-1], label_cols[1:]):
        per_child = df.groupby(c)[p].nunique()
        multi = per_child[per_child > 1]
        det = (f"{df[p].nunique()} nilai '{p}' x {df[c].nunique()} nilai '{c}' berbeda"
               + (f"; nilai '{c}' yang muncul di >1 '{p}': {', '.join(map(str, multi.index[:5]))}"
                  f"{' ...' if len(multi) > 5 else ''}" if len(multi) else ""))
        add("hierarchy", p, c, f"{p} 1..n {c}", df[p].nunique(), None, None, "info", det)

    # R3. Total kolom = jumlah kolom lain dalam grup header yang sama
    groups: dict = {}
    for c in measure_cols:
        groups.setdefault(tuple(paths[c][:-1]), []).append(c)
    for parent, cols in groups.items():
        tot = [c for c in cols if paths[c] and paths[c][-1].casefold() in TOTAL_LABELS]
        oth = [c for c in cols if c not in tot]
        if len(tot) == 1 and len(oth) >= 2:
            lhs = df[tot[0]].astype(float).fillna(0)
            rhs = df[oth].astype(float).fillna(0).sum(axis=1)
            ok = np.isclose(lhs, rhs, rtol=0, atol=1e-6)
            bad = ([" / ".join(map(str, row)) for row in df.loc[~ok, label_cols].head(5).itertuples(index=False)]
                   if not ok.all() else [])
            add("sum_columns", tot[0], ", ".join(oth), f"{tot[0]} = " + " + ".join(oth), n, int(ok.sum()),
                int((~ok).sum()), "ok" if ok.all() else "mismatch",
                None if ok.all() else "Baris tidak cocok: " + "; ".join(bad))

    # R4. Subtotal baris: SUM(baris non-total) = baris Total, per dimensi
    if is_total_row_col in df and df[is_total_row_col].any():
        dim = label_cols[0] if len(label_cols) >= 2 else None
        gb = df.groupby(dim, sort=False) if dim else [(None, df)]
        for c in measure_cols:
            checked = passed = 0
            bad = []
            for d, sub in gb:
                tot_rows = sub[sub[is_total_row_col]]
                if len(tot_rows) != 1 or sub[is_total_row_col].all():
                    continue
                s = float(sub.loc[~sub[is_total_row_col], c].astype(float).fillna(0).sum())
                t = float(tot_rows[c].astype(float).fillna(0).iloc[0])
                checked += 1
                if abs(s - t) <= 1e-6:
                    passed += 1
                else:
                    bad.append(f"{d}: jumlah baris={s:,.0f} vs Total={t:,.0f}")
            if checked:
                add("sum_rows", c, dim or "(tabel)", f"SUM(baris non-total) = baris Total, per {dim or 'tabel'}",
                    checked, passed, checked - passed, "ok" if passed == checked else "mismatch",
                    None if not bad else "; ".join(bad[:5]))

        # R5. Konsistensi grand total antar dimensi
        for c in measure_cols:
            vals = df.loc[df[is_total_row_col], c].dropna().unique()
            add("grand_total_consistency", c, "(semua dimensi)", "baris Total sama di semua dimensi",
                int(df[is_total_row_col].sum()), None, None, "ok" if len(vals) == 1 else "mismatch",
                f"{len(vals)} nilai Total berbeda: {', '.join(f'{v:,.0f}' for v in vals[:5])}")

    # R6. Korelasi antar kolom ukuran (non-total)
    non_total_cols = [c for c in measure_cols if not (paths[c] and paths[c][-1].casefold() in TOTAL_LABELS)]
    if len(non_total_cols) >= 2:
        body = df.loc[~df[is_total_row_col]] if is_total_row_col in df else df
        for a, b in itertools.combinations(non_total_cols, 2):
            r = body[[a, b]].astype(float).corr().iloc[0, 1]
            add("correlation", a, b, "pearson r (baris non-total)", len(body), None, None, "info",
                "Korelasi antar kolom ukuran", None if pd.isna(r) else round(float(r), 4))
    return pd.DataFrame(rel)


# =============================================================================
# PEMROSESAN SHEET DATA
# =============================================================================
@dataclass
class SheetResult:
    sheet: str
    table_wide: str
    table_long: Optional[str]
    wide: pd.DataFrame
    long: Optional[pd.DataFrame]
    dtypes_wide: dict
    dtypes_long: dict
    colmeta: list
    profile: Optional[pd.DataFrame]
    relation: Optional[pd.DataFrame]
    dim_mapping: pd.DataFrame
    recs: list
    sheet_row: dict
    pk_wide: list
    pk_long: list


@dataclass
class ExtractionOptions:
    """Opsi yang mengendalikan pemrosesan sheet data (pengganti argumen CLI)."""
    file: str
    prefix: Optional[str] = None
    const: Dict[str, str] = field(default_factory=dict)
    data_start_row: Optional[int] = None
    profile: bool = True
    relations: bool = True


def process_data_sheet(g: Grid, meta: Metadata, opts: ExtractionOptions, prefix: str, multi: bool) -> SheetResult:
    sheet_slug = slugify(g.name)
    t_wide = f"{prefix}_{sheet_slug}_data" if multi else f"{prefix}_data"
    t_long = f"{prefix}_{sheet_slug}_data_long" if multi else f"{prefix}_data_long"
    filled = g.filled()
    L = detect_layout(g, opts.data_start_row)
    n_cols = L.n_cols
    letters = [get_column_letter(c + 1) for c in range(n_cols)]
    paths = header_paths(filled, L.header_start, L.data_start, n_cols)
    recs: list[dict] = []

    def rec(cat, sev, target, finding, reco):
        recs.append(dict(sheet_name=g.name, category=cat, severity=sev, target=target,
                         finding=finding, recommendation=reco))

    # --- peran kolom: label vs ukuran ---
    roles = []
    for c in range(n_cols):
        nb = [g.values[r][c] for r in L.body_rows if not is_blank(g.values[r][c])]
        nu = sum(1 for v in nb if is_numlike(v))
        roles.append("measure" if nb and nu / len(nb) >= 0.8 else "label")
    names = name_columns(paths, roles, letters)
    label_idx = [c for c in range(n_cols) if roles[c] == "label"]
    meas_idx = [c for c in range(n_cols) if roles[c] == "measure"]
    if not label_idx or not meas_idx:
        raise ValueError(f"Sheet '{g.name}': tidak ditemukan pasangan kolom label & ukuran")
    label_cols = [names[c] for c in label_idx]
    meas_cols = [names[c] for c in meas_idx]
    path_of = {names[c]: paths[c] for c in range(n_cols)}
    letter_of = {names[c]: letters[c] for c in range(n_cols)}

    # --- kolom label: bersihkan, isi sel merge, forward fill ---
    n_merge_filled = n_ws = n_ffill = 0
    lab: dict[str, list] = {}
    for c in label_idx:
        vals = []
        for r in L.body_rows:
            raw, fv = g.values[r][c], filled[r][c]
            if is_blank(raw) and not is_blank(fv):
                n_merge_filled += 1
            t = None if is_blank(fv) else clean_text(fv)
            if not is_blank(raw) and str(raw) != t:
                n_ws += 1
            vals.append(t)
        lab[names[c]] = vals
    lab_df = pd.DataFrame(lab)
    if len(label_cols) > 1:
        before = lab_df[label_cols[:-1]].isna().sum().sum()
        lab_df[label_cols[:-1]] = lab_df[label_cols[:-1]].ffill()
        n_ffill = int(before - lab_df[label_cols[:-1]].isna().sum().sum())

    # --- kolom ukuran: parse angka ---
    meas: dict[str, list] = {}
    dash_cells: dict[str, int] = {}
    nonnum: list[str] = []
    for c in meas_idx:
        vals, nd = [], 0
        for r in L.body_rows:
            p = parse_number(g.values[r][c])
            if p is DASH:
                nd += 1
                p = 0 if DASH_AS_ZERO else None
            elif p is NOT_NUM:
                nonnum.append(f"{letters[c]}{r + 1}={g.values[r][c]!r}")
                p = None
            vals.append(p)
        meas[names[c]] = vals
        dash_cells[names[c]] = nd

    # --- tipe data ---
    dtypes = {c: "text" for c in label_cols}
    for c in meas_cols:
        dtypes[c] = infer_dtype(meas[c])

    # --- konstanta (--const) & lineage ---
    const_cols = list(opts.const.keys())
    now = datetime.now().replace(microsecond=0)
    n = len(L.body_rows)
    wide = {c: to_series(lab_df[c].tolist(), "text") for c in label_cols}
    for c in meas_cols:
        wide[c] = to_series(meas[c], dtypes[c])
    is_total_row = [any(isinstance(v, str) and v.casefold() in TOTAL_LABELS for v in row)
                    for row in lab_df[label_cols].itertuples(index=False)]
    wide["is_total_row"] = pd.array(is_total_row, dtype="boolean")
    for k, v in opts.const.items():
        wide[k] = pd.Series([v] * n, dtype="object")
    wide["_source_row"] = pd.array([r + 1 for r in L.body_rows], dtype="Int64")
    wide["_source_file"] = pd.Series([os.path.basename(opts.file)] * n, dtype="object")
    wide["_loaded_at"] = pd.Series([now] * n)
    wide_df = pd.DataFrame(wide)
    dtypes.update({"is_total_row": "boolean", "_source_row": "bigint", "_source_file": "text",
                   "_loaded_at": "timestamp", **{k: "text" for k in const_cols}})
    wide_df["is_total_row"] = wide_df["is_total_row"].astype(bool)

    # --- pemetaan dimensi <-> metadata (langkah 5) ---
    klas = meta.klasifikasi
    dm_rows = []
    dim_src = label_cols[0]
    for v in pd.unique(wide_df[dim_src].dropna()):
        m, s, how = match_dimension(v, klas)
        cnt = wide_df.loc[wide_df[dim_src] == v]
        dm_rows.append(dict(sheet_name=g.name, source_level=f"baris:{dim_src}", source_value=v,
                            matched_klasifikasi=m, score=s, method=how,
                            n_categories=int((~cnt["is_total_row"]).sum()),
                            has_total_row=bool(cnt["is_total_row"].any())))
    parents = {}
    for c in meas_cols:
        if len(path_of[c]) >= 2:
            parents.setdefault(" > ".join(path_of[c][:-1]), []).append(c)
    for p, cols in parents.items():
        m, s, how = match_dimension(p.split(" > ")[-1], klas)
        dm_rows.append(dict(sheet_name=g.name, source_level="header:kolom", source_value=p,
                            matched_klasifikasi=m, score=s, method=how, n_categories=len(cols),
                            has_total_row=any(path_of[c][-1].casefold() in TOTAL_LABELS for c in cols)))
    dim_mapping = pd.DataFrame(dm_rows)
    mapped_dim = {r["source_value"]: r["matched_klasifikasi"] for r in dm_rows}

    # --- tabel long (unpivot header kolom) ---
    leafs = {c: (path_of[c][-1] if path_of[c] else c) for c in meas_cols}
    grp_of = {c: " > ".join(path_of[c][:-1]) for c in meas_cols}
    distinct_parents = sorted(set(grp_of.values()))
    if len(distinct_parents) == 1 and distinct_parents[0]:
        sub_col, grp_col = slugify(distinct_parents[0].split(" > ")[-1]), None
    else:
        sub_col, grp_col = "kolom", ("kolom_grup" if len(distinct_parents) > 1 else None)
    if sub_col in label_cols:
        sub_col += "_kolom"
    base_cols = label_cols + const_cols + ["is_total_row"]
    frames = []
    for order, c in enumerate(meas_cols):
        f = wide_df[base_cols].copy()
        f[sub_col] = leafs[c]
        if grp_col:
            f[grp_col] = grp_of[c]
        f["is_total_col"] = leafs[c].casefold() in TOTAL_LABELS
        f["nilai"] = pd.to_numeric(wide_df[c].astype(float))
        f["_source_row"] = wide_df["_source_row"]
        f["_col_order"] = order
        frames.append(f)
    long_df = pd.concat(frames, ignore_index=True).sort_values(["_source_row", "_col_order"], kind="stable")
    long_df = long_df.drop(columns="_col_order").reset_index(drop=True)
    nilai_type = "bigint" if all(dtypes[c] == "bigint" for c in meas_cols) else "numeric"
    long_df["nilai"] = to_series(long_df["nilai"].where(long_df["nilai"].notna(), None).tolist(), nilai_type)
    long_df["_source_file"] = os.path.basename(opts.file)
    long_df["_loaded_at"] = now
    ordered = (label_cols + const_cols + [sub_col] + ([grp_col] if grp_col else []) +
               ["is_total_row", "is_total_col", "nilai", "_source_row", "_source_file", "_loaded_at"])
    long_df = long_df[ordered]
    dtypes_long = {**{c: "text" for c in label_cols + const_cols + [sub_col] + ([grp_col] if grp_col else [])},
                   "is_total_row": "boolean", "is_total_col": "boolean", "nilai": nilai_type,
                   "_source_row": "bigint", "_source_file": "text", "_loaded_at": "timestamp"}

    # --- profiling & relasi ---
    prof = rel = None
    if opts.profile:
        extra = {c: {"dash_cells": dash_cells[c]} for c in meas_cols}
        prof = profile_df(g.name, t_wide, wide_df, label_cols + meas_cols + ["is_total_row"], dtypes, extra)
    if opts.relations:
        rel = analyze_relations(g.name, wide_df, label_cols, meas_cols, path_of)
        # kunci kandidat untuk PK
    key_row = rel[rel.relation_type == "candidate_key"] if rel is not None else pd.DataFrame()
    key_cols = ([s.strip() for s in key_row.iloc[0]["left_scope"].split(",")]
                if len(key_row) and key_row.iloc[0]["status"] == "ok" else [])
    pk_wide = (const_cols + key_cols) if key_cols else []
    pk_long = (const_cols + key_cols + [sub_col] + ([grp_col] if grp_col else [])) if key_cols else []

    # --- metadata kolom (langkah 3-5) ---
    colmeta: list[dict] = []

    def cm(table, col, ordinal, role, src, hpath, dtype, desc, defn, match, score, dsrc):
        colmeta.append(dict(sheet_name=g.name, table_name=table, column_name=col, ordinal_position=ordinal,
                            column_role=role, source_col=src, header_path=hpath, inferred_type=dtype,
                            pg_type=PG_NAMES[dtype], description=desc, definition=defn,
                            metadata_match=match, match_score=score, description_source=dsrc))

    n_dim_ok = sum(1 for r in dm_rows if r["source_level"].startswith("baris") and r["matched_klasifikasi"])
    n_dim = sum(1 for r in dm_rows if r["source_level"].startswith("baris"))
    for i, c in enumerate(wide_df.columns, 1):
        d = dtypes[c]
        if c in label_cols:
            k = label_cols.index(c)
            grouped = len(label_cols) > 1 and paths[label_idx[0]] == paths[label_idx[min(k, len(label_idx) - 1)]]
            role = ("dimensi" if k == 0 else "kategori") if len(label_cols) > 1 else "label"
            if k == 0 and len(label_cols) > 1:
                desc = ((f"Nama dimensi/klasifikasi baris. {n_dim_ok}/{n_dim} dimensi berhasil dipetakan ke "
                         f"daftar Klasifikasi di metadata (lihat tabel dim_mapping).") if klas
                        else "Nama dimensi/klasifikasi baris.")
            elif len(label_cols) > 1:
                desc = f"Kategori/nilai dari dimensi pada kolom '{label_cols[0]}'."
            else:
                desc = f"Kolom label baris '{' > '.join(path_of[c])}'."
            cm(t_wide, c, i, role, letter_of[c], " > ".join(path_of[c]), d, desc, None, None, None, "inferred")
        elif c in meas_cols:
            info = indicator_info(meta, path_of[c])
            desc, defn, dsrc = describe_measure(info, path_of[c])
            cm(t_wide, c, i, "ukuran", letter_of[c], " > ".join(path_of[c]), d, desc, defn,
               info.get("nama"), None, dsrc)
        elif c == "is_total_row":
            cm(t_wide, c, i, "flag_turunan", None, None, d,
               "TRUE bila baris adalah subtotal ('Total') suatu dimensi. Jangan dijumlahkan bersama baris lain.",
               None, None, None, "derived")
        elif c in const_cols:
            cm(t_wide, c, i, "konstanta", None, None, d, f"Konstanta dari argumen --const ({c}).",
               None, None, None, "cli")
        else:
            cm(t_wide, c, i, "lineage", None, None, d,
               {"_source_row": "Nomor baris Excel asal (1-based).", "_source_file": "Nama file sumber.",
                "_loaded_at": "Waktu proses load."}[c], None, None, None, "derived")
    info_long = indicator_info(meta, [])
    for i, c in enumerate(long_df.columns, 1):
        d = dtypes_long[c]
        if c in label_cols or c in const_cols:
            src = next(x for x in colmeta if x["table_name"] == t_wide and x["column_name"] == c)
            cm(t_long, c, i, src["column_role"], src["source_col"], src["header_path"], d,
               src["description"], None, None, None, src["description_source"])
        elif c == sub_col:
            pn = distinct_parents[0] if len(distinct_parents) == 1 else None
            m = mapped_dim.get(pn) if pn else None
            cm(t_long, c, i, "dimensi_kolom", None, pn, d,
               f"Nilai header kolom Excel '{pn or 'kolom'}' (hasil unpivot)."
               + (f" Dipetakan ke klasifikasi metadata: {m}." if m else ""), None, m, None, "derived")
        elif c == grp_col:
            cm(t_long, c, i, "dimensi_kolom", None, None, d, "Header induk kolom Excel (hasil unpivot).",
               None, None, None, "derived")
        elif c == "is_total_col":
            cm(t_long, c, i, "flag_turunan", None, None, d,
               "TRUE bila header kolom adalah Total. Hindari double count.", None, None, None, "derived")
        elif c == "is_total_row":
            cm(t_long, c, i, "flag_turunan", None, None, d, "TRUE bila baris adalah subtotal ('Total').",
               None, None, None, "derived")
        elif c == "nilai":
            desc, defn, dsrc = describe_measure(info_long, [])
            cm(t_long, c, i, "ukuran", None, None, d, desc or "Nilai ukuran.", defn, info_long.get("nama"), None, dsrc)
        else:
            src = next(x for x in colmeta if x["table_name"] == t_wide and x["column_name"] == c)
            cm(t_long, c, i, src["column_role"], None, None, d, src["description"], None, None, None, "derived")

    # --- rekomendasi (langkah 9) ---
    nh = L.data_start - L.header_start
    if nh > 1:
        rec("struktur", "info", g.name,
            f"Header {nh} level (baris {L.header_start + 1}-{L.data_start}) digabung jadi satu nama kolom.",
            f"Gunakan tabel {t_long} (unpivot header kolom) untuk analisis/dbt; tabel {t_wide} untuk pengecekan.")
    if n_merge_filled or n_ffill:
        rec("struktur", "info", ", ".join(label_cols[:-1] or label_cols),
            f"{n_merge_filled} sel merge diisi nilai induknya, {n_ffill} sel diisi forward-fill.",
            "Tidak perlu tindakan; jejak baris asli ada di kolom _source_row.")
    n_tr, n_tc = int(wide_df["is_total_row"].sum()), sum(1 for c in meas_cols if leafs[c].casefold() in TOTAL_LABELS)
    if n_tr or n_tc:
        rec("kualitas", "warn", "is_total_row / is_total_col",
            f"Ada {n_tr} baris Total dan {n_tc} kolom Total tercampur dengan data detail.",
            "Selalu filter is_total_row = false (dan is_total_col = false) saat melakukan SUM/AVG agar tidak double count.")
    nd = sum(dash_cells.values())
    if nd:
        zero = "0" if DASH_AS_ZERO else "NULL"
        rec("kualitas", "warn", ", ".join(c for c in meas_cols if dash_cells[c]),
            f"{nd} sel berisi '-' (teks) diubah menjadi {zero}.",
            "Konvensi BPS: '-' = tidak ada/nol. Lihat tabel relation (sum_columns/sum_rows): bila status ok, asumsi 0 terbukti konsisten.")
    if n_ws:
        rec("kualitas", "info", ", ".join(label_cols), f"{n_ws} sel label memiliki spasi berlebih (leading/trailing/ganda) dan dirapikan.",
            "Tidak perlu tindakan.")
    if nonnum:
        rec("kualitas", "warn", ", ".join(meas_cols), f"{len(nonnum)} sel di kolom ukuran bukan angka dan diisi NULL: {'; '.join(nonnum[:5])}",
            "Periksa sel tersebut di Excel sumber.")
    if L.dropped_rows:
        rec("kualitas", "warn", g.name,
            f"{len(L.dropped_rows)} baris di tengah tabel tanpa angka dibuang (baris Excel: {[r + 1 for r in L.dropped_rows][:10]}).",
            "Pastikan baris tersebut memang judul/sub-judul, bukan data.")
    prefixed = {c: int(wide_df[c].dropna().str.match(r"^\s*\d+[.\s]").sum()) for c in label_cols[-1:]}
    for c, k in prefixed.items():
        if k:
            rec("model", "info", c, f"{k} nilai '{c}' diawali nomor/kode (mis. '1. Berusaha sendiri').",
                "Di layer staging dbt pertimbangkan memisah kode & label (hati-hati: '0 Jam (Sementara Tidak Bekerja)' bukan kode).")
    if not klas:
        rec("metadata", "warn", "metadata",
            "Tidak ada sheet/daftar Klasifikasi metadata; arti kolom hanya disimpulkan dari header.",
            "Lengkapi deskripsi kolom secara manual (COMMENT ON COLUMN) atau sediakan sheet 'Standar Data & Metadata'.")
    unmatched = [r["source_value"] for r in dm_rows if not r["matched_klasifikasi"]] if klas else []
    if unmatched:
        rec("metadata", "warn", "dim_mapping", f"Dimensi tanpa padanan di metadata: {unmatched}.",
            "Tambahkan istilah ke SYNONYMS di bagian konfigurasi script.")
    low = [f"{r['source_value']} -> {r['matched_klasifikasi']} ({r['score']})" for r in dm_rows
           if r["matched_klasifikasi"] and r["score"] < 0.85]
    if low:
        rec("metadata", "info", "dim_mapping", f"Pemetaan berskor <0.85 (verifikasi manual): {low}.",
            "Cek tabel dim_mapping dan koreksi bila salah.")
    used = {r["matched_klasifikasi"] for r in dm_rows if r["matched_klasifikasi"]}
    missing = [k for k in klas if k not in used]
    if missing:
        rec("metadata", "info", "klasifikasi", f"Klasifikasi di metadata yang tidak menjadi kolom/dimensi di data: {missing}.",
            "Bila file ini mewakili satu nilai (mis. satu provinsi), simpan sebagai kolom konstanta via --const.")
    texts = " \n".join(str(v) for v in list(meta.dataset["value"].dropna()) +
                       (list(meta.indicator.astype(str).values.ravel()) if meta.indicator is not None else []))
    prov = sorted(set(re.findall(r"[Pp]rovinsi +([A-Z][\w]*(?: +[A-Z][\w]*)*)", texts)))
    per = sorted(set(re.findall(rf"[Pp]eriode +((?:{MONTHS})\b +\d{{4}})", texts)))
    ctx = [f"{k}" for k in ("provinsi", "periode") if k not in const_cols and ((k == "provinsi" and prov) or (k == "periode" and per))]
    if ctx:
        hint = " ".join(([f"--const provinsi='{prov[0]}'"] if "provinsi" in ctx else []) +
                        ([f"--const periode='{per[0]}'"] if "periode" in ctx else []))
        rec("model", "warn", "konteks", f"Tabel data tidak memuat {' & '.join(ctx)}, padahal metadata menyebutkannya.",
            f"Tambahkan kolom konteks agar beberapa file bisa digabung tanpa tertukar: {hint}")
    if pk_long:
        rec("model", "info", t_long, f"Kunci unik terdeteksi: ({', '.join(pk_long)}).",
            "Dipakai sebagai kandidat PRIMARY KEY saat data dimuat ke tabel tujuan.")
    rec("model", "info", t_wide, "Nama dimensi + kategori berulang di setiap baris.",
        "Untuk data warehouse, bentuk dim_karakteristik (karakteristik, kategori, surrogate key) lalu fact(dim_key, kolom_dimensi, nilai).")
    if rel is not None:
        for _, rr in rel[rel.status == "mismatch"].iterrows():
            rec("relasi", "high", f"{rr.left_scope}", f"{rr.relation_type}: {rr.expression} -> {rr.failed or ''} gagal. {rr.detail or ''}",
                "Verifikasi ke publikasi/sumber asli; mungkin ada kategori terpotong, dibulatkan, atau salah ketik.")
        ok = rel[rel.status == "ok"]
        if len(ok):
            rec("relasi", "info", "relation", f"{len(ok)} pemeriksaan relasi lolos ({', '.join(sorted(set(ok.relation_type)))}).",
                "Data konsisten secara internal; aman dipakai.")

    sheet_row = dict(sheet_name=g.name, role="data", n_rows=g.n_rows, n_cols=g.n_cols,
                     n_merged_ranges=len(g.merged), header_first_row=L.header_start + 1,
                     header_last_row=L.data_start, data_first_row=L.data_start + 1, data_last_row=L.data_end + 1,
                     label_columns=", ".join(label_cols), measure_columns=", ".join(meas_cols),
                     notes="; ".join(clean_text(g.values[r][0]) if not is_blank(g.values[r][0]) else
                                     clean_text(next(v for v in g.values[r] if not is_blank(v)))
                                     for r in L.title_rows + L.footer_rows))
    # catatan footer ke meta_dataset
    for r in L.footer_rows + L.title_rows:
        txt = next((clean_text(v) for v in g.values[r] if not is_blank(v)), None)
        if txt:
            kv = kv_from_text(txt)
            meta.dataset.loc[len(meta.dataset)] = [g.name, f"Catatan sheet '{g.name}'", kv[0] if kv else "Catatan",
                                                  kv[1] if kv else txt, f"A{r + 1}", len(meta.dataset) + 1]
    return SheetResult(g.name, t_wide, t_long, wide_df, long_df, dtypes, dtypes_long, colmeta, prof, rel,
                       dim_mapping, recs, sheet_row, pk_wide, pk_long)


# Tipe Postgres hanya dipakai sebagai label metadata (kolom pg_type), bukan untuk load.
PG_NAMES = {"bigint": "BIGINT", "numeric": "NUMERIC", "boolean": "BOOLEAN", "date": "DATE",
            "timestamp": "TIMESTAMP", "text": "TEXT"}

# =============================================================================
# OUTPUT PILIHAN
# =============================================================================
OUTPUT_CHOICES = ["data", "data_long", "dim_mapping", "meta_sheet", "meta_dataset", "meta_indicator",
                  "meta_column", "profile", "relation", "recommendation"]
OUTPUT_ALIASES = {"relations": "relation", "long": "data_long", "wide": "data", "rekomendasi": "recommendation"}


def parse_outputs(values, parser=None) -> list[str]:
    """['data,data_long', '_profile'] -> ['data', 'data_long', 'profile'] (urutan tetap, tanpa duplikat)."""
    out: list[str] = []
    for chunk in values:
        for tok in re.split(r"[,\s]+", chunk):
            name = tok.strip().lower().lstrip("_")
            if not name:
                continue
            if name == "all":
                return list(OUTPUT_CHOICES)
            name = OUTPUT_ALIASES.get(name, name)
            if name not in OUTPUT_CHOICES:
                msg = f"output '{tok}' tidak dikenal. Pilihan: {', '.join(OUTPUT_CHOICES)}, all"
                if parser is not None:
                    parser.error(msg)
                raise ValueError(msg)
            if name not in out:
                out.append(name)
    return out


# =============================================================================
# API REUSABLE: BACA WORKBOOK -> TABEL (DICT DATAFRAME)
# =============================================================================
@dataclass
class ExtractionResult:
    """Hasil otomasi membaca Excel: tabel siap CSV untuk normalisasi & loaders."""
    tables: Dict[str, pd.DataFrame]
    prefix: str
    recommendation: pd.DataFrame
    sheets: List[SheetResult]
    metadata: Metadata


def extract_workbook(
    path: str,
    *,
    prefix: Optional[str] = None,
    const: Optional[Dict[str, str]] = None,
    data_start_row: Optional[int] = None,
    profile: bool = True,
    relations: bool = True,
    outputs: Optional[List[str]] = None,
) -> ExtractionResult:
    """Baca workbook Excel dan kembalikan tabel terstruktur (DataFrame).

    Args:
        path: path file .xlsx.
        prefix: prefix nama tabel; default dari nama sheet data.
        const: kolom konstanta tambahan, mis. {"provinsi": "Gorontalo"}.
        data_start_row: paksa baris awal data (1-based) bila deteksi salah.
        profile: hitung tabel profiling.
        relations: hitung tabel relasi (dipakai untuk deteksi kunci).
        outputs: batasi tabel keluaran (lihat OUTPUT_CHOICES). None = semua.

    Returns:
        ExtractionResult dengan ``.tables`` berisi ``{nama_tabel: DataFrame}``.
    """
    const_dict = {slugify(k): v for k, v in (const or {}).items()}
    opts = ExtractionOptions(file=path, prefix=prefix, const=const_dict,
                             data_start_row=data_start_row, profile=profile, relations=relations)

    # 1. baca semua sheet
    grids = load_grids(path)
    roles = detect_roles(grids)
    data_grids = [g for g in grids if roles[g.name] == "data"]
    meta_grids = [g for g in grids if roles[g.name] == "metadata"]
    resolved_prefix = prefix or slugify(data_grids[0].name, 40)

    # 5-6. metadata
    meta = Metadata(pd.DataFrame(columns=["sheet_name", "section", "key", "value", "source_cell", "ordinal"]),
                    None, {}, {}, [], {})
    sheet_rows = []
    for mg in meta_grids:
        meta, info = parse_metadata(mg)
        sheet_rows.append(dict(sheet_name=mg.name, role="metadata", n_rows=mg.n_rows, n_cols=mg.n_cols,
                               n_merged_ranges=len(mg.merged),
                               header_first_row=info.get("header_first"), header_last_row=info.get("header_last"),
                               data_first_row=info.get("data_first"), data_last_row=info.get("data_last"),
                               label_columns=None, measure_columns=None,
                               notes=f"{info.get('n_kv', 0)} pasangan key-value; {info.get('n_indicators', 0)} indikator; "
                                     f"{info.get('n_klasifikasi', 0)} klasifikasi"))
        break                                           # satu sheet metadata dipakai

    # 2-4, 7-9. sheet data
    results = [process_data_sheet(g, meta, opts, resolved_prefix, len(data_grids) > 1) for g in data_grids]
    sheet_rows = [r.sheet_row for r in results] + sheet_rows

    # --- susun tabel keluaran ---
    tables: Dict[str, pd.DataFrame] = {}
    colmeta_all = [cm for r in results for cm in r.colmeta]
    if meta.indicator is not None:                      # kolom meta_indicator (nama asli header disimpan)
        for i, c in enumerate(meta.indicator.columns, 1):
            dt = infer_dtype(meta.indicator[c].dropna().tolist())
            colmeta_all.append(dict(sheet_name=meta_grids[0].name, table_name=f"{resolved_prefix}_meta_indicator",
                                    column_name=c, ordinal_position=i, column_role="metadata",
                                    source_col=meta.indicator_letters[c], header_path=" > ".join(meta.indicator_paths[c]),
                                    inferred_type=dt, pg_type=PG_NAMES[dt],
                                    description=" > ".join(meta.indicator_paths[c]), definition=None,
                                    metadata_match=None, match_score=None, description_source="header sheet metadata"))
    meta_col_df = pd.DataFrame(colmeta_all)
    meta_col_df["match_score"] = pd.to_numeric(meta_col_df["match_score"])
    meta_col_df["ordinal_position"] = pd.array(meta_col_df["ordinal_position"], dtype="Int64")

    for r in results:
        tables[r.table_wide] = r.wide
        if r.long is not None:
            tables[r.table_long] = r.long
    tables[f"{resolved_prefix}_dim_mapping"] = pd.concat([r.dim_mapping for r in results], ignore_index=True)
    tables[f"{resolved_prefix}_meta_sheet"] = pd.DataFrame(sheet_rows)
    tables[f"{resolved_prefix}_meta_dataset"] = meta.dataset
    if meta.indicator is not None:
        tables[f"{resolved_prefix}_meta_indicator"] = meta.indicator
    tables[f"{resolved_prefix}_meta_column"] = meta_col_df
    if profile:
        tables[f"{resolved_prefix}_profile"] = pd.concat([r.profile for r in results], ignore_index=True)
    if relations:
        tables[f"{resolved_prefix}_relation"] = pd.concat([r.relation for r in results], ignore_index=True)
    rec_df = pd.DataFrame([x for r in results for x in r.recs])
    sev = {"high": 0, "warn": 1, "info": 2}
    if not rec_df.empty:
        rec_df = rec_df.sort_values("severity", key=lambda s: s.map(sev), kind="stable").reset_index(drop=True)
        rec_df.insert(0, "no", range(1, len(rec_df) + 1))
    tables[f"{resolved_prefix}_recommendation"] = rec_df

    # outputs: sisakan hanya tabel yang diminta
    if outputs is not None:
        keep = set()
        for name in outputs:
            if name == "data":
                keep |= {r.table_wide for r in results}
            elif name == "data_long":
                keep |= {r.table_long for r in results}
            else:
                keep.add(f"{resolved_prefix}_{name}")
        tables = {k: v for k, v in tables.items() if k in keep}

    return ExtractionResult(tables=tables, prefix=resolved_prefix, recommendation=rec_df,
                            sheets=results, metadata=meta)


def write_csvs(tables: Dict[str, pd.DataFrame], out_dir: str) -> List[str]:
    """Tulis setiap tabel ke CSV (UTF-8 BOM, tanpa index). Kembalikan path yang ditulis."""
    os.makedirs(out_dir, exist_ok=True)
    written: List[str] = []
    for name, df in tables.items():
        path = os.path.join(out_dir, f"{name}.csv")
        df.to_csv(path, index=False, encoding="utf-8-sig")
        written.append(path)
    return written


# =============================================================================
# CLI (selalu keluaran CSV)
# =============================================================================
def parse_args():
    p = argparse.ArgumentParser(description="Excel (data + metadata) -> tabel CSV terstruktur",
                                formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    p.add_argument("file", help="path file .xlsx")
    p.add_argument("--prefix", help="prefix nama tabel (default: dari nama sheet data)")
    p.add_argument("--const", action="append", default=[], metavar="KEY=VALUE",
                   help="kolom konstanta tambahan, boleh berulang (mis. --const provinsi=Gorontalo)")
    p.add_argument("--data-start-row", type=int, help="paksa baris awal data (1-based) bila deteksi salah")
    p.add_argument("--output", action="extend", nargs="+", metavar="NAMA",
                   help="tabel yang dikeluarkan; pisahkan dengan koma atau spasi. Pilihan: "
                        + ", ".join(OUTPUT_CHOICES) + ", all. Default: data. "
                        "Contoh: --output data,data_long,recommendation")
    p.add_argument("--out-dir", default="./out", help="direktori keluaran CSV (default: ./out)")
    a = p.parse_args()
    a.const_dict = {}
    for item in a.const:
        if "=" not in item:
            p.error(f"--const harus KEY=VALUE, bukan '{item}'")
        k, v = item.split("=", 1)
        a.const_dict[slugify(k)] = v
    a.outputs = parse_outputs(a.output or ["data"], p)
    return a


def _print_summary(result: ExtractionResult) -> None:
    print(f"Sheet ditemukan: {[(r.sheet, 'data') for r in result.sheets]}")
    for r in result.sheets:
        print(f"\nSheet '{r.sheet}': header baris {r.sheet_row['header_first_row']}-{r.sheet_row['header_last_row']}, "
              f"data baris {r.sheet_row['data_first_row']}-{r.sheet_row['data_last_row']}")
        print(r.wide.dtypes.astype(str).to_string())
    print("\nREKOMENDASI:")
    if result.recommendation.empty:
        print("(tidak ada)")
    else:
        print(result.recommendation[["no", "severity", "category", "target", "finding"]]
              .to_string(index=False, max_colwidth=110))
    print()


def main() -> None:
    args = parse_args()
    outputs = args.outputs
    result = extract_workbook(
        args.file,
        prefix=args.prefix,
        const=args.const_dict,
        data_start_row=args.data_start_row,
        profile="profile" in outputs,
        relations=True,
        outputs=outputs,
    )
    _print_summary(result)
    written = write_csvs(result.tables, args.out_dir)
    for path in written:
        print(f"  [CSV] {path}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:                            # noqa: BLE001
        print(f"ERROR: {exc.__class__.__name__}: {exc}", file=sys.stderr)
        raise
