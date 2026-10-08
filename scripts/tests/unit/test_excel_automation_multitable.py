"""Tests for multi-table-per-sheet extraction and the sheet_name column."""

from openpyxl import Workbook

from ingestion.core.extractors.excel_automation import (
    extract_workbook,
    split_by_markers,
)


def _make_workbook(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Tabel 1. Jumlah A"])
    ws.append(["Provinsi", "Jumlah"])
    ws.append(["Aceh", 10])
    ws.append(["Bali", 20])
    ws.append([])
    ws.append(["Tabel 2. Jumlah B"])
    ws.append(["Provinsi", "Lain"])
    ws.append(["Aceh", 1])
    ws.append(["Bali", 2])
    path = tmp_path / "multi.xlsx"
    wb.save(path)
    return str(path)


def test_sheet_name_column_present_with_default(tmp_path):
    path = _make_workbook(tmp_path)
    result = extract_workbook(path, prefix="xl", profile=False, relations=False, outputs=["data"])

    assert len(result.tables) == 1
    (df,) = result.tables.values()
    assert "_sheet_name" in df.columns
    assert (df["_sheet_name"] == "Data").all()


def test_multi_table_split_by_markers(tmp_path):
    path = _make_workbook(tmp_path)
    result = extract_workbook(
        path,
        prefix="xl",
        profile=False,
        relations=False,
        outputs=["data"],
        table_markers=["Tabel 1", "Tabel 2"],
    )

    names = sorted(result.tables)
    assert names == ["xl_data_tabel_1_data", "xl_data_tabel_2_data"]

    t1 = result.tables["xl_data_tabel_1_data"]
    t2 = result.tables["xl_data_tabel_2_data"]

    assert list(t1["provinsi"]) == ["Aceh", "Bali"]
    assert list(t1["jumlah"]) == [10, 20]
    assert list(t2["lain"]) == [1, 2]
    for df in (t1, t2):
        assert (df["_sheet_name"] == "Data").all()


def test_split_by_markers_returns_whole_grid_without_marker(tmp_path):
    from ingestion.core.extractors.excel_automation import load_grids

    path = _make_workbook(tmp_path)
    (grid,) = load_grids(path)
    units = split_by_markers(grid, ["TidakAda"])
    assert len(units) == 1
    label, sub = units[0]
    assert label is None
    assert sub.n_rows == grid.n_rows


def test_non_data_marker_segment_is_skipped_not_fatal(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Tabel 1. Data A"])
    ws.append(["Provinsi", "Jumlah"])
    ws.append(["Aceh", 10])
    ws.append([])
    ws.append(["Tabel 2. Referensi"])          # text-only table, no numeric column
    ws.append(["Kode", "Nama"])
    ws.append(["A", "Alpha"])
    path = tmp_path / "skip.xlsx"
    wb.save(path)

    result = extract_workbook(
        str(path), prefix="xl", profile=False, relations=False,
        outputs=["data"], table_markers=["Tabel 1", "Tabel 2"],
    )

    assert list(result.tables) == ["xl_data_tabel_1_data"]
    assert any("dilewati" in str(f) for f in result.recommendation["finding"])


def test_duplicate_marker_gets_unique_table_names(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Tabel 1. A"])
    ws.append(["Provinsi", "Jumlah"])
    ws.append(["Aceh", 1])
    ws.append([])
    ws.append(["Tabel 1. B"])                  # same marker repeated
    ws.append(["Provinsi", "Jumlah"])
    ws.append(["Bali", 2])
    path = tmp_path / "dup.xlsx"
    wb.save(path)

    result = extract_workbook(
        str(path), prefix="xl", profile=False, relations=False,
        outputs=["data"], table_markers=["Tabel 1"],
    )

    names = sorted(result.tables)
    assert names == ["xl_data_tabel_1_2_data", "xl_data_tabel_1_data"]

